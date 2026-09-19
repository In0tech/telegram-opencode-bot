from __future__ import annotations

import asyncio
import base64
import json
import re
import time
import uuid
from dataclasses import dataclass
from pathlib import Path

from openai import OpenAI
from pptx import Presentation
from pptx.util import Inches, Pt

from config import Settings


class CreativeError(RuntimeError):
    pass


def _safe_filename(value: str, fallback: str) -> str:
    value = re.sub(r'[^\w.-]+', '_', value, flags=re.UNICODE).strip('._')
    return value[:80] or fallback


def _extract_json(text: str) -> dict:
    text = text.strip()
    fence = chr(96) * 3
    if text.startswith(fence):
        lines = text.splitlines()
        if lines:
            lines = lines[1:]
        if lines and lines[-1].strip() == fence:
            lines = lines[:-1]
        text = '\n'.join(lines).strip()

    try:
        value = json.loads(text)
        if isinstance(value, dict):
            return value
    except json.JSONDecodeError:
        pass

    start = text.find('{')
    end = text.rfind('}')
    if start >= 0 and end > start:
        value = json.loads(text[start:end + 1])
        if isinstance(value, dict):
            return value

    raise CreativeError('Модель вернула некорректный JSON для презентации.')


@dataclass
class TarotCard:
    name: str
    orientation: str


class CreativeService:
    def __init__(self, settings: Settings):
        self.settings = settings
        self.output_dir = settings.state_dir / 'generated'
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def _client(self) -> OpenAI:
        if not self.settings.openai_api_key:
            raise CreativeError(
                'OPENAI_API_KEY не настроен. Добавьте его в .env и перезапустите бота.'
            )
        return OpenAI(api_key=self.settings.openai_api_key)

    async def draft_email(self, brief: str) -> str:
        return await asyncio.to_thread(self._draft_email_sync, brief)

    def _draft_email_sync(self, brief: str) -> str:
        client = self._client()
        response = client.responses.create(
            model=self.settings.openai_text_model,
            input=(
                'Напиши готовое письмо по заданию ниже. По умолчанию используй русский язык, '
                'если пользователь явно не просит другой. Верни только готовый текст: сначала '
                'строка "Тема: ...", затем тело письма. Не выдумывай факты, имена или адреса.\n\n'
                f'Задание:\n{brief}'
            ),
        )
        return response.output_text.strip()

    async def create_presentation(self, brief: str, slides: int = 8) -> Path:
        return await asyncio.to_thread(self._create_presentation_sync, brief, slides)

    def _create_presentation_sync(self, brief: str, slides: int) -> Path:
        client = self._client()
        slides = min(max(slides, 3), 20)
        response = client.responses.create(
            model=self.settings.openai_text_model,
            input=(
                'Создай структуру профессиональной презентации. Верни СТРОГО JSON без markdown '
                'по схеме: {"title":"...", "subtitle":"...", "slides":[{"title":"...",'
                '"bullets":["...","..."]}]}. Заголовочный слайд НЕ включай в slides. '
                f'Нужно {slides - 1} содержательных слайдов. На каждом 3-6 коротких пунктов. '
                'Не выдумывай конкретные цифры и источники, если их нет в задании.\n\n'
                f'Задание пользователя:\n{brief}'
            ),
        )

        deck = _extract_json(response.output_text)
        title = str(deck.get('title') or 'Презентация')
        subtitle = str(deck.get('subtitle') or '')
        slide_items = deck.get('slides')
        if not isinstance(slide_items, list) or not slide_items:
            raise CreativeError('Не удалось получить структуру слайдов.')

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        title_slide = prs.slides.add_slide(prs.slide_layouts[0])
        title_slide.shapes.title.text = title
        if len(title_slide.placeholders) > 1:
            title_slide.placeholders[1].text = subtitle

        for item in slide_items[:slides - 1]:
            if not isinstance(item, dict):
                continue
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = str(item.get('title') or 'Слайд')
            tf = slide.placeholders[1].text_frame
            tf.clear()
            bullets = item.get('bullets') or []
            if not isinstance(bullets, list):
                bullets = [str(bullets)]
            for idx, bullet in enumerate(bullets[:6]):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                p.text = str(bullet)
                p.level = 0
                p.font.size = Pt(22)

        name = _safe_filename(title, 'presentation')
        path = self.output_dir / f'{name}_{uuid.uuid4().hex[:8]}.pptx'
        prs.save(path)
        return path

    async def create_image(self, prompt: str) -> Path:
        return await asyncio.to_thread(self._create_image_sync, prompt)

    def _create_image_sync(self, prompt: str) -> Path:
        client = self._client()
        result = client.images.generate(
            model=self.settings.openai_image_model,
            prompt=prompt,
            size=self.settings.openai_image_size,
            quality=self.settings.openai_image_quality,
            output_format='png',
        )
        if not result.data or not result.data[0].b64_json:
            raise CreativeError('Image API не вернул изображение.')

        data = base64.b64decode(result.data[0].b64_json)
        path = self.output_dir / f'image_{uuid.uuid4().hex[:10]}.png'
        path.write_bytes(data)
        return path

    async def create_video(self, prompt: str) -> Path:
        return await asyncio.to_thread(self._create_video_sync, prompt)

    def _create_video_sync(self, prompt: str) -> Path:
        client = self._client()

        try:
            video = client.videos.create(
                model=self.settings.openai_video_model,
                prompt=prompt,
                seconds=str(self.settings.openai_video_seconds),
                size=self.settings.openai_video_size,
            )
        except Exception as exc:
            raise CreativeError(f'Не удалось запустить генерацию видео: {exc}') from exc

        deadline = time.monotonic() + self.settings.video_timeout_seconds
        while getattr(video, 'status', None) in {'queued', 'in_progress'}:
            if time.monotonic() >= deadline:
                raise CreativeError(
                    f'Видео ещё не готово после {self.settings.video_timeout_seconds} секунд. '
                    f'Job ID: {video.id}'
                )
            time.sleep(5)
            video = client.videos.retrieve(video.id)

        if getattr(video, 'status', None) != 'completed':
            err = getattr(getattr(video, 'error', None), 'message', None)
            raise CreativeError(
                err or f'Генерация видео завершилась со статусом {video.status}'
            )

        content = client.videos.download_content(video.id, variant='video')
        path = self.output_dir / f'video_{uuid.uuid4().hex[:10]}.mp4'
        if hasattr(content, 'write_to_file'):
            content.write_to_file(str(path))
        else:
            path.write_bytes(content.read())
        return path

    async def interpret_tarot(self, question: str, cards: list[TarotCard]) -> str:
        return await asyncio.to_thread(self._interpret_tarot_sync, question, cards)

    def _interpret_tarot_sync(self, question: str, cards: list[TarotCard]) -> str:
        client = self._client()
        spread = '\n'.join(
            f'{idx + 1}. {card.name} — {card.orientation}'
            for idx, card in enumerate(cards)
        )
        response = client.responses.create(
            model=self.settings.openai_text_model,
            input=(
                'Сделай вдумчивую интерпретацию расклада Таро на русском языке. '
                'Подавай это как символический инструмент для размышления, а не как гарантированное '
                'предсказание будущего. Структура: значение каждой карты, общая картина, практический '
                'вопрос для размышления. Не утверждай неизбежность событий.\n\n'
                f'Вопрос: {question or "Общий расклад"}\n'
                f'Карты:\n{spread}'
            ),
        )
        return response.output_text.strip()
